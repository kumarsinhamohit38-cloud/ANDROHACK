"""Extract, chunk, and search files uploaded to the current app session."""

import csv
import hashlib
import io
import re
from html.parser import HTMLParser
from pathlib import Path

import numpy as np


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        if data.strip():
            self.parts.append(data.strip())


def extract_document_text(filename: str, content: bytes) -> str:
    """Return readable text from supported document types."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        for table in document.tables:
            for row in table.rows:
                text += "\n" + " | ".join(cell.text.strip() for cell in row.cells)
    elif suffix == ".csv":
        decoded = content.decode("utf-8-sig", errors="replace")
        text = "\n".join(
            " | ".join(cell.strip() for cell in row if cell.strip())
            for row in csv.reader(io.StringIO(decoded))
        )
    else:
        decoded = content.decode("utf-8-sig", errors="replace")
        if suffix in {".html", ".htm"}:
            parser = _HTMLTextExtractor()
            parser.feed(decoded)
            text = "\n".join(parser.parts)
        else:
            text = decoded

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if not text:
        if suffix == ".pdf":
            raise ValueError("No selectable text was found. This may be a scanned PDF; OCR is not enabled yet.")
        raise ValueError("No readable text was found in this file.")
    return text


def split_into_chunks(text: str, max_chars: int = 1000, overlap: int = 140) -> list[str]:
    """Split extracted text into overlapping chunks that keep local context."""
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            boundary = max(text.rfind(". ", start, end), text.rfind("; ", start, end), text.rfind(" ", start, end - 80))
            if boundary > start + max_chars // 2:
                end = boundary + (1 if text[boundary] in ".;" else 0)
        chunk = text[start:end].strip()
        if chunk and (not chunks or chunk != chunks[-1]):
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def prepare_uploaded_document(filename: str, content: bytes, embedding_model) -> dict:
    """Extract and embed one uploaded file for fast session-only retrieval."""
    text = extract_document_text(filename, content)
    chunks = split_into_chunks(text)
    if not chunks:
        raise ValueError("No searchable passages could be created from this file.")
    embeddings = embedding_model.encode_documents(chunks, show_progress=False)
    return {
        "fingerprint": hashlib.sha256(content).hexdigest(),
        "filename": Path(filename).name,
        "chunks": chunks,
        "embeddings": embeddings,
    }


def search_uploaded_documents(query: str, documents: dict, embedding_model, top_k: int = 5) -> list[dict]:
    """Search uploaded document chunks by cosine similarity."""
    if not query.strip() or not documents:
        return []

    query_vector = embedding_model.encode_query(query)
    candidates = []
    for document in documents.values():
        embeddings = document["embeddings"]
        scores = np.asarray(embeddings) @ query_vector
        for chunk_index, (text, score) in enumerate(zip(document["chunks"], scores)):
            candidates.append({
                "document_id": f"upload_{document['fingerprint'][:12]}_{chunk_index + 1}",
                "text": text,
                "score": float(score),
                "rank": 0,
                "retrieval_method": "uploaded document",
                "metadata": {
                    "source": "uploaded_document",
                    "dataset": document["filename"],
                    "topic": f"{document['filename']} · passage {chunk_index + 1}",
                },
            })

    candidates.sort(key=lambda item: item["score"], reverse=True)
    selected = candidates[:max(1, top_k)]
    for rank, result in enumerate(selected, start=1):
        result["rank"] = rank
    return selected
